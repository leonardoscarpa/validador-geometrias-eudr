import datetime

TIER_ICON = {"Low": "🟢", "Medium": "🟡", "High": "🔴"}


def generate_markdown(scores, warnings, input_filename):
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    lines = []
    lines.append("# Supplier risk scorecard: EUDR (Art. 10)")
    lines.append("")
    lines.append(f"**Input file:** `{input_filename}`  ")
    lines.append(f"**Generated:** {now}  ")
    lines.append("**Regulatory basis:** Regulation (EU) 2023/1115, Art. 10 (risk assessment criteria). "
                  "Country classification per Implementing Regulation (EU) 2025/1093.")
    lines.append("")

    by_tier = {"High": [], "Medium": [], "Low": []}
    for s in scores:
        by_tier[s.tier].append(s)

    lines.append("## Executive summary")
    lines.append("")
    lines.append("| Metric | Value |")
    lines.append("|---|---|")
    lines.append(f"| Suppliers evaluated | {len(scores)} |")
    lines.append(f"| 🔴 High risk | {len(by_tier['High'])} |")
    lines.append(f"| 🟡 Medium risk | {len(by_tier['Medium'])} |")
    lines.append(f"| 🟢 Low risk | {len(by_tier['Low'])} |")
    lines.append("")
    lines.append(
        "**Terminology note:** \"Low / Medium / High\" here is the risk level of the "
        "*supplier* computed by this tool. Do not confuse it with \"Low / Standard / High\", "
        "which is the *official country classification* under Regulation 2025/1093. The country is "
        "one input to the supplier score and does not determine it on its own."
    )
    lines.append("")

    for tier in ("High", "Medium", "Low"):
        group = sorted(by_tier[tier], key=lambda s: -s.total_points)
        if not group:
            continue
        lines.append(f"## {TIER_ICON[tier]} {tier} risk ({len(group)})")
        lines.append("")
        for s in group:
            lines.append(f"### {s.producer_name}: {s.total_points} points")
            lines.append(f"Parcels: {', '.join(s.parcels)} · Total area: {s.total_area_ha} ha · "
                          f"Country(ies): {', '.join(sorted(s.countries)) or '(no data)'}")
            lines.append("")
            for f in s.factors:
                sign = "+" if f.points >= 0 else ""
                lines.append(f"- **{sign}{f.points}** [{f.article}] {f.label}: {f.detail}")
            lines.append("")

    if warnings:
        lines.append("## ⚠️ Data loading warnings")
        lines.append("")
        for w in warnings:
            lines.append(f"- {w}")
        lines.append("")

    return "\n".join(lines)


def generate_excel(scores, warnings, input_filename, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()

    ws_f = wb.active
    ws_f.title = "Factors"
    ws_f.append(["Supplier", "Points", "Article", "Factor", "Detail"])
    for s in scores:
        for f in s.factors:
            ws_f.append([s.producer_name, f.points, f.article, f.label, f.detail])

    ws_s = wb.create_sheet("Suppliers")
    ws_s.append(["Supplier", "Parcels", "Area (ha)", "Countries", "Total score", "Level"])
    for i, s in enumerate(scores, start=2):
        ws_s.append([s.producer_name, len(s.parcels), s.total_area_ha,
                     ", ".join(sorted(s.countries)), s.total_points, s.tier])

    ws_r = wb.create_sheet("Summary", 0)
    ws_r.append(["Supplier risk scorecard: EUDR (Art. 10)"])
    ws_r.append(["Input file", input_filename])
    ws_r.append(["Generated", datetime.datetime.now().strftime("%Y-%m-%d %H:%M")])
    ws_r.append([])
    ws_r.append(["Metric", "Value"])
    last_s = len(scores) + 1
    ws_r.append(["Suppliers evaluated", f"=COUNTA(Suppliers!A2:A{last_s})" if scores else 0])
    ws_r.append(["High risk", f'=COUNTIF(Suppliers!F2:F{last_s},"High")' if scores else 0])
    ws_r.append(["Medium risk", f'=COUNTIF(Suppliers!F2:F{last_s},"Medium")' if scores else 0])
    ws_r.append(["Low risk", f'=COUNTIF(Suppliers!F2:F{last_s},"Low")' if scores else 0])
    ws_r.append([])
    ws_r.append(["Note: \"Low/Medium/High\" is the risk level of the SUPPLIER computed by this "
                  "tool, different from \"Low/Standard/High\" (OFFICIAL country classification, "
                  "Regulation 2025/1093), which is only one of the inputs to the score."])

    header_fill = PatternFill(start_color="2F5233", end_color="2F5233", fill_type="solid")
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    for ws in (ws_f, ws_s):
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
        for col_idx, width in enumerate([26, 10, 16, 34, 70][:ws.max_column], start=1):
            ws.column_dimensions[chr(64 + col_idx)].width = width
    ws_r["A1"].font = Font(name="Arial", bold=True, size=14)
    ws_r["A5"].font = header_font; ws_r["A5"].fill = header_fill
    ws_r["B5"].font = header_font; ws_r["B5"].fill = header_fill
    ws_r.column_dimensions["A"].width = 30
    ws_r.column_dimensions["B"].width = 70
    for ws in (ws_r, ws_s, ws_f):
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    wb.save(path)
