"""
Generates the quality report in three formats: Markdown, PDF and Excel.
The target audience is a compliance team working in SAP/Excel, not remote
sensing specialists, so the language avoids geospatial jargon and every
finding cites its regulatory/technical basis.
"""

import datetime

from .models import ERROR, WARNING, INFO

SEVERITY_LABEL = {
    ERROR: "Blocking error",
    WARNING: "Warning",
    INFO: "Conversion note",
}
SEVERITY_LABEL_PLURAL = {
    ERROR: "Blocking errors",
    WARNING: "Warnings",
    INFO: "Conversion notes",
}
SEVERITY_ICON = {ERROR: "🔴", WARNING: "🟡", INFO: "ℹ️"}


def build_summary(gdf, findings, failed_ids=None):
    failed_ids = list(failed_ids or [])
    bad_ids = set()
    n_error = n_warning = n_info = 0
    for f in findings:
        if f.severity == ERROR:
            n_error += 1
            if f.parcel_id:
                for pid in str(f.parcel_id).split(" / "):
                    bad_ids.add(pid)
        elif f.severity == WARNING:
            n_warning += 1
        elif f.severity == INFO:
            n_info += 1

    # Counted per RECORD: GeoDataFrame rows (parcels that got a geometry)
    # PLUS the ParcelIDs that failed so early in ingestion that they never
    # formed a geometry (missing lat/lon, prohibited geometry, etc.). Those
    # are also "parcels evaluated with error" and must not vanish from the
    # count.
    n_gdf = len(gdf) if not gdf.empty else 0
    total_parcels = n_gdf + len(failed_ids)
    bad_rows_in_gdf = int(gdf["ParcelID"].astype(str).isin(bad_ids).sum()) if n_gdf else 0
    bad_parcels = bad_rows_in_gdf + len(failed_ids)
    ok_parcels = total_parcels - bad_parcels

    return {
        "total_parcels": total_parcels,
        "ok_parcels": max(ok_parcels, 0),
        "bad_parcels": bad_parcels,
        "n_error": n_error,
        "n_warning": n_warning,
        "n_info": n_info,
    }


def _sorted_findings(findings):
    return sorted(findings, key=lambda f: f.sort_key())


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------

def generate_markdown(gdf, findings, input_filename, commodity="Wood (Southern Cone)", failed_ids=None):
    s = build_summary(gdf, findings, failed_ids)
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    lines = []
    lines.append("# Quality report: EUDR supplier geometry validator")
    lines.append("")
    lines.append(f"**Input file:** `{input_filename}`  ")
    lines.append(f"**Generated:** {now}  ")
    lines.append(f"**Commodity:** {commodity}  ")
    lines.append("**Reference format:** EUDR GeoJSON File Description v1.5 "
                  "(Implementing Regulation (EU) 2024/3084)")
    lines.append("")
    lines.append("## Executive summary")
    lines.append("")
    lines.append("| Metric | Value |")
    lines.append("|---|---|")
    lines.append(f"| Parcels evaluated | {s['total_parcels']} |")
    lines.append(f"| Parcels ready to submit as-is | {s['ok_parcels']} |")
    lines.append(f"| Parcels with blocking error | {s['bad_parcels']} |")
    lines.append(f"| Blocking errors (total) | {s['n_error']} |")
    lines.append(f"| Warnings (non-blocking) | {s['n_warning']} |")
    lines.append(f"| Conversion notes | {s['n_info']} |")
    lines.append("")
    lines.append(
        "**What \"ready\" means:** the parcel has no finding of "
        "severity *Blocking error*. The output GeoJSON (`output.geojson`) "
        "contains only ready parcels. Parcels with a blocking error "
        "are left out of the output file and must be corrected at the source before "
        "attempting submission again."
    )
    lines.append("")

    sorted_findings = _sorted_findings(findings)
    for sev in (ERROR, WARNING, INFO):
        subset = [f for f in sorted_findings if f.severity == sev]
        if not subset:
            continue
        lines.append(f"## {SEVERITY_ICON[sev]} {SEVERITY_LABEL_PLURAL[sev]} ({len(subset)})")
        lines.append("")
        for f in subset:
            pid_txt = f" Parcel `{f.parcel_id}`" if f.parcel_id else " Whole file"
            lines.append(f"- **[{f.code}]**{pid_txt}: {f.message}")
            lines.append(f"  *Basis: {f.reference}*")
        lines.append("")

    if not findings:
        lines.append("No findings detected: all parcels pass the applied "
                      "checks. ✅")
        lines.append("")

    return "\n".join(lines)


# ---------------------------------------------------------------------------
# PDF (reportlab)
# ---------------------------------------------------------------------------

def generate_pdf(gdf, findings, input_filename, path, commodity="Wood (Southern Cone)", failed_ids=None):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak,
    )

    s = build_summary(gdf, findings, failed_ids)
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="Small", parent=styles["Normal"], fontSize=8, leading=10))
    styles.add(ParagraphStyle(name="FindingMsg", parent=styles["Normal"], fontSize=9, leading=12))

    doc = SimpleDocTemplate(path, pagesize=A4,
                             leftMargin=18 * mm, rightMargin=18 * mm,
                             topMargin=16 * mm, bottomMargin=16 * mm)
    story = []

    story.append(Paragraph("Quality report: EUDR supplier geometry validator",
                            styles["Title"]))
    story.append(Spacer(1, 4 * mm))
    meta_tbl = Table([
        ["Input file", input_filename],
        ["Generated", now],
        ["Commodity", commodity],
        ["Reference format", "EUDR GeoJSON File Description v1.5 (Regulation 2024/3084)"],
    ], colWidths=[45 * mm, 120 * mm])
    meta_tbl.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("TEXTCOLOR", (0, 0), (0, -1), colors.grey),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(meta_tbl)
    story.append(Spacer(1, 6 * mm))

    story.append(Paragraph("Executive summary", styles["Heading2"]))
    summary_data = [
        ["Metric", "Value"],
        ["Parcels evaluated", str(s["total_parcels"])],
        ["Parcels ready to submit as-is", str(s["ok_parcels"])],
        ["Parcels with blocking error", str(s["bad_parcels"])],
        ["Blocking errors (total)", str(s["n_error"])],
        ["Warnings (non-blocking)", str(s["n_warning"])],
        ["Conversion notes", str(s["n_info"])],
    ]
    sum_tbl = Table(summary_data, colWidths=[100 * mm, 40 * mm])
    sum_tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2f5233")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f2f2")]),
    ]))
    story.append(sum_tbl)
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph(
        "<b>What \"ready\" means:</b> the parcel has no finding of "
        "severity Blocking error. The output GeoJSON contains only the "
        "ready parcels.", styles["Small"]))
    story.append(PageBreak())

    sorted_findings = _sorted_findings(findings)
    sev_colors = {
        ERROR: colors.HexColor("#c0392b"),
        WARNING: colors.HexColor("#b7950b"),
        INFO: colors.HexColor("#2874a6"),
    }
    for sev in (ERROR, WARNING, INFO):
        subset = [f for f in sorted_findings if f.severity == sev]
        if not subset:
            continue
        story.append(Paragraph(f"{SEVERITY_LABEL_PLURAL[sev]} ({len(subset)})",
                                ParagraphStyle(name=f"H_{sev}", parent=styles["Heading2"],
                                               textColor=sev_colors[sev])))
        story.append(Spacer(1, 2 * mm))
        rows = [["Parcel", "Code", "Detail"]]
        for f in subset:
            pid_txt = f.parcel_id if f.parcel_id else "Whole file"
            pid_p = Paragraph(str(pid_txt), styles["FindingMsg"])
            code_p = Paragraph(f.code, styles["FindingMsg"])
            detail = Paragraph(f"{f.message}<br/><font size=7 color='grey'>{f.reference}</font>",
                                styles["FindingMsg"])
            rows.append([pid_p, code_p, detail])
        t = Table(rows, colWidths=[36 * mm, 36 * mm, 98 * mm], repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8e8e8")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 9),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(t)
        story.append(Spacer(1, 6 * mm))

    if not findings:
        story.append(Paragraph("No findings detected: all parcels "
                                "pass the applied checks.", styles["Normal"]))

    doc.build(story)


# ---------------------------------------------------------------------------
# Excel (openpyxl, with formulas so the summary recalculates)
# ---------------------------------------------------------------------------

def generate_excel(gdf, findings, areas_ha, input_filename, path, commodity="Wood (Southern Cone)",
                    failed_ids=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    failed_ids = list(failed_ids or [])
    wb = Workbook()

    # ---- Findings sheet (base data) ----
    ws_f = wb.active
    ws_f.title = "Findings"
    headers = ["Severity", "ParcelID", "Code", "Message", "Regulatory reference"]
    ws_f.append(headers)
    for f in _sorted_findings(findings):
        ws_f.append([f.severity, f.parcel_id or "(whole file)", f.code, f.message, f.reference])

    # ---- Parcels sheet (one row per parcel, with a Status formula) ----
    ws_p = wb.create_sheet("Parcels")
    ws_p.append(["ParcelID", "Geometry type", "Computed area (ha)",
                 "Declared area (ha)", "ProducerCountry", "Status"])
    n_parcels = 0
    if not gdf.empty:
        for i, (idx, row) in enumerate(gdf.iterrows(), start=2):
            n_parcels += 1
            gtype = row["geometry"].geom_type if row["geometry"] is not None else ""
            area_calc = round(float(areas_ha.loc[idx]), 3) if idx in areas_ha.index else None
            declared = row.get("Area")
            try:
                declared = float(declared)
            except (TypeError, ValueError):
                declared = None
            ws_p.append([
                row["ParcelID"], gtype, area_calc, declared, row.get("ProducerCountry"),
                # COUNTIFS with wildcards: supports the PARCELS_OVERLAP case,
                # whose ParcelID in Findings is combined ("A / B") because it
                # involves two parcels at once.
                f'=IF(COUNTIFS(Findings!A:A,"ERROR",Findings!B:B,"*"&A{i}&"*")>0,"Has error","Ready")'
            ])
    # Parcels that failed so early in ingestion that they never got a
    # geometry (missing lat/lon, prohibited geometry, etc.) are still
    # listed, with empty geometric fields, so they do not silently
    # disappear from the report.
    for pid in failed_ids:
        n_parcels += 1
        i = n_parcels + 1
        ws_p.append([
            pid, "(no valid geometry)", None, None, None,
            f'=IF(COUNTIFS(Findings!A:A,"ERROR",Findings!B:B,"*"&A{i}&"*")>0,"Has error","Ready")'
        ])

    # ---- Summary sheet (formulas over the other two sheets) ----
    ws_r = wb.create_sheet("Summary", 0)
    ws_r.append(["Quality report: EUDR supplier geometry validator"])
    ws_r.append(["Input file", input_filename])
    ws_r.append(["Generated", datetime.datetime.now().strftime("%Y-%m-%d %H:%M")])
    ws_r.append(["Commodity", commodity])
    ws_r.append(["Reference format", "EUDR GeoJSON File Description v1.5 (Regulation 2024/3084)"])
    ws_r.append([])
    ws_r.append(["Metric", "Value"])
    last_p = n_parcels + 1
    last_f = len(findings) + 1
    ws_r.append(["Parcels evaluated", f"=COUNTA(Parcels!A2:A{last_p})" if n_parcels else 0])
    ws_r.append(["Parcels ready to submit as-is",
                  f'=COUNTIF(Parcels!F2:F{last_p},"Ready")' if n_parcels else 0])
    ws_r.append(["Parcels with blocking error",
                  f'=COUNTIF(Parcels!F2:F{last_p},"Has error")' if n_parcels else 0])
    ws_r.append(["Blocking errors (total)",
                  f'=COUNTIF(Findings!A2:A{last_f},"ERROR")' if findings else 0])
    ws_r.append(["Warnings (non-blocking)",
                  f'=COUNTIF(Findings!A2:A{last_f},"WARNING")' if findings else 0])
    ws_r.append(["Conversion notes",
                  f'=COUNTIF(Findings!A2:A{last_f},"INFO")' if findings else 0])
    ws_r.append([])
    ws_r.append(["What \"ready\" means: the parcel has no finding of "
                  "severity Blocking error. The output GeoJSON contains "
                  "only the ready parcels."])

    # ---- style ----
    header_fill = PatternFill(start_color="2F5233", end_color="2F5233", fill_type="solid")
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    title_font = Font(name="Arial", bold=True, size=14)
    normal_font = Font(name="Arial", size=10)

    for ws in (ws_f, ws_p):
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = normal_font
                cell.alignment = Alignment(vertical="top", wrap_text=(ws is ws_f and cell.column == 4))
        for col_idx in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 22
        if ws is ws_f:
            ws.column_dimensions["D"].width = 70
            ws.column_dimensions["E"].width = 55

    ws_r["A1"].font = title_font
    for r in range(2, 6):
        ws_r.cell(row=r, column=1).font = Font(name="Arial", bold=True, size=10)
        ws_r.cell(row=r, column=2).font = normal_font
    ws_r["A7"].font = header_font
    ws_r["B7"].font = header_font
    ws_r["A7"].fill = header_fill
    ws_r["B7"].fill = header_fill
    for r in range(8, 14):
        ws_r.cell(row=r, column=1).font = normal_font
        ws_r.cell(row=r, column=2).font = Font(name="Arial", size=10, bold=True)
    ws_r.column_dimensions["A"].width = 42
    ws_r.column_dimensions["B"].width = 70

    # print setup: landscape, fit to 1 page wide
    # (in case someone prints it or exports it to PDF, besides opening it in Excel)
    for ws in (ws_r, ws_p, ws_f):
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = False

    wb.save(path)
